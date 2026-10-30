"""Tableau de suivi des photos de Karl dans Wikipédia (docs/wiki-route.md).

Outil de session, à relancer après chaque ajout : il lit les photos déjà dans un article
(vitrine/wikipedia.csv) et les pistes (pistes.csv, à côté de ce fichier), puis écrit
suivi-wikipedia.xlsx et consigne-rassemblement.txt, la consigne que Karl colle dans Claude
sur son Mac pour rassembler les originaux des photos prêtes avant de les verser sur
Commons. Il demande openpyxl (python3 -m pip install openpyxl), absent du site lui-même.

    python3 docs/wiki-suivi/fabriquer.py
"""

import csv
from datetime import date
from pathlib import Path
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ICI = Path(__file__).resolve().parent
RACINE = ICI.parent.parent

LANGUES = {"fr": "français", "en": "anglais", "es": "espagnol", "eu": "basque", "ast": "asturien",
           "it": "italien", "de": "allemand", "gl": "galicien", "ca": "catalan", "oc": "occitan",
           "br": "breton", "pt": "portugais"}
# Ordre d'affichage et couleur de chaque statut.
STATUTS = {
    "Prête": "FFF2CC",           # case préparée : verser la photo sur Commons, puis l'article
    "À vérifier": "DDEBF7",      # sujet identifié, article à vérifier par la session
    "À identifier": "EDE7F6",    # Karl doit dire ce que montre la photo
    "En attente": "F2F2F2",      # pas de place pour l'instant
    "Dans l'article": "E2EFDA",  # en ligne, et dans le panneau du site
    "Écartée": "FBE5D6",         # raison dans la colonne Remarques
}
POLICE = "Arial"


def lire(chemin):
    with open(chemin, encoding="utf-8", newline="") as f:
        return [{k: (v or "").strip() for k, v in ligne.items() if k} for ligne in csv.DictReader(f)]


def lien_article(wiki, article, section=""):
    base = f"https://{wiki}.wikipedia.org"
    if section != "":
        return f"{base}/w/index.php?title={quote(article.replace(' ', '_'))}&action=edit&section={section}"
    return f"{base}/wiki/{quote(article.replace(' ', '_'))}"


def lien_commons(fichier):
    return "https://commons.wikimedia.org/wiki/File:" + quote(fichier.replace(" ", "_"))


def lignes():
    titres = {l["photo"]: l["titre_fr"] for l in lire(RACINE / "vitrine" / "donnees" / "textes-fr.csv")}
    faites = []
    for w in lire(RACINE / "vitrine" / "wikipedia.csv"):
        faites.append({
            "photo": w["photo"], "sujet": titres.get(w["photo"]) or w.get("titre_fr") or w["fichier"],
            "statut": "Dans l'article", "wiki": w["wiki"], "article": w["article"],
            "partie": w["partie"] or "infobox", "legende": w["legende"], "fichier": w["fichier"],
            "date": w["ajoutee_le"], "section": "",
            "remarques": "" if w["photo"] else "Seulement sur Commons : le panneau mène à sa page Commons."})
    faites.sort(key=lambda l: l["date"], reverse=True)
    pistes = [dict(p, date="") for p in lire(ICI / "pistes.csv")]
    ordre = list(STATUTS)
    pistes.sort(key=lambda p: ordre.index(p["statut"]) if p["statut"] in ordre else len(ordre))
    avant = [p for p in pistes if p["statut"] != "Écartée"]
    return avant + faites + [p for p in pistes if p["statut"] == "Écartée"]


def consigne(pretes):
    liste = "\n".join(f"- {p['photo']} → {p['fichier']}" for p in pretes)
    return f"""Je verse mes photos sur Wikimedia Commons. Rassemble les fichiers ORIGINAUX des photos ci-dessous dans le dossier ~/Downloads/Commons à verser/ (crée-le s'il n'existe pas).

Les photos (numéro Pexels → nom à donner au fichier) :
{liste}

Méthode :
1. Trouve où mon application Telepex garde ses données (par exemple dans ~/Library/Application Support/ ou ~/Library/Containers/, ou dans ses réglages) et lis-les sans rien y modifier. Pour chaque numéro Pexels, retrouve le fichier original sur mon Mac : son chemin ou son nom de fichier d'appareil. Si Telepex ne garde pas ce lien, cherche dans mes dossiers de photos d'après la date de publication et le titre que Telepex donne pour ce numéro, et montre-moi les fichiers candidats avant de choisir.
2. COPIE chaque original dans le dossier (ne déplace rien, ne supprime rien, ne modifie aucun fichier d'origine), en le renommant comme indiqué. Garde chaque fichier intact : pas de recadrage, pas de compression, métadonnées EXIF conservées.
3. Ne télécharge rien depuis Pexels et n'envoie rien sur Internet.
4. À la fin, donne-moi un tableau : numéro Pexels, nom du fichier copié, chemin de l'original, taille en pixels, date de prise de vue (exiftool -DateTimeOriginal), coordonnées GPS s'il y en a. Si l'année de prise de vue diffère de celle du nom, dis-le-moi. Signale chaque photo introuvable, avec les pistes essayées.
"""


def feuille_suivi(classeur, toutes):
    f = classeur.active
    f.title = "Suivi"
    entetes = ["N° Pexels", "Photo", "Statut", "Wikipédia", "Article", "Partie", "Légende",
               "Fichier sur Commons", "Ajoutée le", "Lien de modification", "Remarques"]
    largeurs = [12, 42, 14, 11, 30, 18, 42, 46, 12, 22, 60]
    f.append(entetes)
    for l in toutes:
        f.append([int(l["photo"]) if l["photo"].isdigit() else "", l["sujet"], l["statut"],
                  LANGUES.get(l["wiki"], l["wiki"]), l["article"], l["partie"], l["legende"], l["fichier"],
                  l["date"], "Modifier la partie" if l["statut"] == "Prête" and l["section"] != "" else "",
                  l["remarques"]])
        n = f.max_row
        if l["photo"].isdigit():
            f.cell(n, 1).hyperlink = f"https://www.pexels.com/photo/{l['photo']}/"
        if l["article"] and l["wiki"]:
            f.cell(n, 5).hyperlink = lien_article(l["wiki"], l["article"])
        if l["fichier"] and l["statut"] in ("Dans l'article", "En attente"):
            f.cell(n, 8).hyperlink = lien_commons(l["fichier"])
        if f.cell(n, 10).value:
            f.cell(n, 10).hyperlink = lien_article(l["wiki"], l["article"], l["section"])
        fond = PatternFill("solid", fgColor=STATUTS.get(l["statut"], "FFFFFF"))
        for c in range(1, len(entetes) + 1):
            cellule = f.cell(n, c)
            cellule.fill = fond
            cellule.alignment = Alignment(vertical="top", wrap_text=True)
            cellule.font = Font(name=POLICE, size=10, color="0563C1" if cellule.hyperlink else "000000",
                                underline="single" if cellule.hyperlink else None)
    trait = Side(style="thin", color="BFBFBF")
    for c, largeur in enumerate(largeurs, 1):
        f.column_dimensions[get_column_letter(c)].width = largeur
        tete = f.cell(1, c)
        tete.font = Font(name=POLICE, size=10, bold=True, color="FFFFFF")
        tete.fill = PatternFill("solid", fgColor="404040")
        tete.alignment = Alignment(vertical="center", wrap_text=True)
        tete.border = Border(bottom=trait)
    f.freeze_panes = "C2"
    f.auto_filter.ref = f"A1:{get_column_letter(len(entetes))}{f.max_row}"
    f.page_setup.orientation = "landscape"
    f.page_setup.fitToWidth, f.page_setup.fitToHeight = 1, 0
    f.sheet_properties.pageSetUpPr.fitToPage = True
    f.print_title_rows = "1:1"


def feuille_bilan(classeur, nombre):
    f = classeur.create_sheet("Bilan")
    f.append(["Statut", "Photos", "Ce que cela veut dire"])
    sens = {
        "Prête": "Case préparée : la photo est à verser sur Commons, puis à ajouter à l'article.",
        "À vérifier": "Sujet identifié : la session vérifie l'article et prépare la case.",
        "À identifier": "Karl doit dire ce que montre la photo.",
        "En attente": "Pas de bonne place pour l'instant.",
        "Dans l'article": "En ligne sur Wikipédia, et dans le panneau du site photo.",
        "Écartée": "Raison dans la colonne Remarques de l'onglet Suivi.",
    }
    for statut, texte in sens.items():
        f.append([statut, f'=COUNTIF(Suivi!$C$2:$C${nombre + 1},A{f.max_row + 1})', texte])
    f.append(["Total", f"=SUM(B2:B{f.max_row})", "Lignes de l'onglet Suivi (une par photo et par article)."])
    f.append([])
    f.append([f"Mis à jour le {date.today().strftime('%d/%m/%Y')} à partir de vitrine/wikipedia.csv et "
              "docs/wiki-suivi/pistes.csv (dépôt PexelsWillwonder)."])
    for ligne in f.iter_rows():
        for cellule in ligne:
            cellule.font = Font(name=POLICE, size=10, bold=cellule.row == 1 or cellule.value == "Total")
            cellule.alignment = Alignment(vertical="top", wrap_text=True)
            if cellule.column == 1 and cellule.value in STATUTS:
                cellule.fill = PatternFill("solid", fgColor=STATUTS[cellule.value])
    for c, largeur in zip("ABC", (16, 10, 80)):
        f.column_dimensions[c].width = largeur


def feuille_rassembler(classeur, pretes, texte):
    f = classeur.create_sheet("À rassembler")
    f.append(["N° Pexels", "Nom du fichier à donner", "Wikipédia", "Article"])
    for p in pretes:
        f.append([int(p["photo"]), p["fichier"], LANGUES.get(p["wiki"], p["wiki"]), p["article"]])
        f.cell(f.max_row, 1).hyperlink = f"https://www.pexels.com/photo/{p['photo']}/"
    f.append([])
    f.append(["Consigne à coller dans Claude, sur le Mac, quand toutes les photos de la liste sont prêtes :"])
    debut = f.max_row + 1
    f.append([texte])
    f.merge_cells(start_row=debut, start_column=1, end_row=debut, end_column=4)
    f.row_dimensions[debut].height = 15 * (texte.count("\n") + 12)
    for ligne in f.iter_rows():
        for cellule in ligne:
            cellule.font = Font(name=POLICE, size=10, bold=cellule.row == 1,
                                color="0563C1" if cellule.hyperlink else "000000")
            cellule.alignment = Alignment(vertical="top", wrap_text=True)
    for c, largeur in zip("ABCD", (12, 60, 12, 40)):
        f.column_dimensions[c].width = largeur


def main():
    toutes = lignes()
    pretes = [l for l in toutes if l["statut"] == "Prête"]
    texte = consigne(pretes)
    classeur = Workbook()
    feuille_suivi(classeur, toutes)
    feuille_bilan(classeur, len(toutes))
    feuille_rassembler(classeur, pretes, texte)
    classeur.calculation.fullCalcOnLoad = True
    classeur.save(ICI / "suivi-wikipedia.xlsx")
    (ICI / "consigne-rassemblement.txt").write_text(texte, encoding="utf-8")
    print(f"{len(toutes)} lignes, dont {len(pretes)} photos prêtes : suivi-wikipedia.xlsx et consigne-rassemblement.txt écrits.")


if __name__ == "__main__":
    main()
